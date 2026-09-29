"""Tests for the Excel journal. Every file is created in pytest's tmp_path."""

from __future__ import annotations

import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import journal
from journal import HEADER, Record, SaveStatus, append_record, clean_name

WHEN = datetime(2026, 9, 29, 10, 15, 30)


def read_rows(path: Path) -> list[tuple[Any, ...]]:
    """Return all row values of the first sheet."""
    workbook = load_workbook(path)
    try:
        return [tuple(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def make_existing_journal(path: Path, rows: list[tuple[Any, ...]]) -> None:
    """Create a workbook the way Excel or an earlier run would have left it."""
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def fail_once(exc: Exception, real: Any) -> Any:
    """Return a stand-in for *real* that raises *exc* on the first call only."""
    calls = {"count": 0}

    def wrapper(*args: Any, **kwargs: Any) -> Any:
        calls["count"] += 1
        if calls["count"] == 1:
            raise exc
        return real(*args, **kwargs)

    return wrapper


# --- file creation and appending -------------------------------------------------


def test_new_file_is_created_with_header(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    assert not path.exists()

    result = append_record(path, Record("Иванов", "Иван", 7, WHEN))

    assert result.ok
    rows = read_rows(path)
    assert rows[0] == ("Last name", "First name", "Номер билета", "Дата и время")
    assert rows[0] == HEADER
    assert rows[1:] == [("Иванов", "Иван", 7, WHEN)]


def test_datetime_column_uses_required_format(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    append_record(path, Record("Иванов", "Иван", 7, WHEN.replace(microsecond=654321)))

    cell = load_workbook(path).worksheets[0]["D2"]
    assert cell.number_format == "yyyy-mm-dd hh:mm:ss"
    assert cell.value == WHEN  # whole seconds, as shown in Excel
    assert cell.value.strftime("%Y-%m-%d %H:%M:%S") == "2026-09-29 10:15:30"


def test_append_adds_exactly_one_row_and_keeps_existing_rows(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    existing = [
        HEADER,
        ("Петров", "Пётр", 3, datetime(2026, 9, 1, 9, 0, 0)),
        ("Сидорова", "Анна", 20, "added by hand"),
    ]
    make_existing_journal(path, existing)
    before = read_rows(path)

    result = append_record(path, Record("Иванов", "Иван", 7, WHEN))

    assert result.ok
    after = read_rows(path)
    assert len(after) == len(before) + 1
    assert after[: len(before)] == before
    assert after[-1] == ("Иванов", "Иван", 7, WHEN)


def test_two_consecutive_appends_are_stored_in_order(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    first = Record("Иванов", "Иван", 7, WHEN)
    second = Record("Петрова", "Мария", 12, WHEN.replace(minute=16))

    assert append_record(path, first).ok
    assert append_record(path, second).ok

    assert read_rows(path) == [HEADER, first.as_row(), second.as_row()]


def test_empty_existing_sheet_gets_header_first(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    make_existing_journal(path, [])

    assert append_record(path, Record("Иванов", "Иван", 7, WHEN)).ok

    assert read_rows(path) == [HEADER, ("Иванов", "Иван", 7, WHEN)]


def test_rows_with_only_formatting_do_not_create_a_gap(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER)
    sheet.append(("Петров", "Пётр", 3, WHEN))
    sheet["A10"].font = Font(bold=True)  # styled but empty, e.g. left over from Excel
    workbook.save(path)

    assert append_record(path, Record("Иванов", "Иван", 7, WHEN)).ok

    sheet = load_workbook(path).worksheets[0]
    assert sheet["A3"].value == "Иванов"
    assert sheet["A10"].value is None
    assert sheet["A10"].font.bold


def test_names_that_look_like_formulas_are_stored_as_text(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    assert append_record(path, Record("=1+1", "+Иван", 7, WHEN)).ok

    sheet = load_workbook(path).worksheets[0]
    assert (sheet["A2"].value, sheet["A2"].data_type) == ("=1+1", "s")
    assert (sheet["B2"].value, sheet["B2"].data_type) == ("+Иван", "s")


# --- trimming and empty-input validation -----------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Иванов", "Иванов"),
        ("  Иванов  ", "Иванов"),
        ("\tПётр \n", "Пётр"),
        ("Анна Мария", "Анна Мария"),
        ("Ива\x1bнов", "Иванов"),  # control characters are dropped
    ],
)
def test_clean_name_trims_surrounding_whitespace(raw: str, expected: str) -> None:
    assert clean_name(raw) == expected


@pytest.mark.parametrize("raw", ["", " ", "     ", "\t", " \t \n ", "\x1b"])
def test_clean_name_rejects_blank_input(raw: str) -> None:
    assert clean_name(raw) is None


def test_record_trims_names_before_saving(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    record = Record("  Иванов ", "\tИван  ", 7, WHEN)

    assert (record.last_name, record.first_name) == ("Иванов", "Иван")
    append_record(path, record)
    assert read_rows(path)[1][:2] == ("Иванов", "Иван")


@pytest.mark.parametrize(("last", "first"), [("", "Иван"), ("Иванов", "   ")])
def test_record_refuses_blank_names(last: str, first: str) -> None:
    with pytest.raises(ValueError):
        Record(last, first, 7, WHEN)


# --- error handling --------------------------------------------------------------


def test_permission_error_during_save_is_reported_not_raised(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "journal.xlsx"
    assert append_record(path, Record("Петров", "Пётр", 3, WHEN)).ok
    original = path.read_bytes()

    def locked(*_args: Any) -> None:
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(journal.os, "replace", locked)
    result = append_record(path, Record("Иванов", "Иван", 7, WHEN))

    assert result.status is SaveStatus.LOCKED
    assert not result.ok
    assert isinstance(result.error, PermissionError)
    assert path.read_bytes() == original  # earlier records are untouched
    assert list(tmp_path.iterdir()) == [path]  # temp file cleaned up


def test_permission_error_while_opening_is_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "journal.xlsx"
    make_existing_journal(path, [HEADER])

    def locked(*_args: Any, **_kwargs: Any) -> None:
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(journal, "load_workbook", locked)

    assert append_record(path, Record("Иванов", "Иван", 7, WHEN)).status is SaveStatus.LOCKED


def test_same_record_is_saved_once_after_retry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "journal.xlsx"
    record = Record("Иванов", "Иван", 7, WHEN)
    monkeypatch.setattr(journal.os, "replace", fail_once(PermissionError(13, "locked"), journal.os.replace))

    assert append_record(path, record).status is SaveStatus.LOCKED
    assert append_record(path, record).ok

    assert read_rows(path) == [HEADER, record.as_row()]


@pytest.mark.skipif(sys.platform != "win32", reason="Windows locks files that are open")
def test_file_held_open_by_another_program_is_reported_as_locked(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    assert append_record(path, Record("Петров", "Пётр", 3, WHEN)).ok

    with open(path, "rb"):  # like Excel, keeps the file from being replaced
        result = append_record(path, Record("Иванов", "Иван", 7, WHEN))
    assert result.status is SaveStatus.LOCKED

    assert append_record(path, Record("Иванов", "Иван", 7, WHEN)).ok
    assert len(read_rows(path)) == 3


def test_corrupted_file_is_reported_and_left_untouched(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    path.write_bytes(b"this is not an Excel workbook")

    result = append_record(path, Record("Иванов", "Иван", 7, WHEN))

    assert result.status is SaveStatus.CORRUPTED
    assert result.error is not None
    assert path.read_bytes() == b"this is not an Excel workbook"


def test_zip_without_workbook_is_reported_as_corrupted(tmp_path: Path) -> None:
    path = tmp_path / "journal.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "no workbook here")

    assert append_record(path, Record("Иванов", "Иван", 7, WHEN)).status is SaveStatus.CORRUPTED


def test_other_os_errors_are_reported(tmp_path: Path) -> None:
    path = tmp_path / "no-such-dir" / "journal.xlsx"

    result = append_record(path, Record("Иванов", "Иван", 7, WHEN))

    assert result.status is SaveStatus.IO_ERROR
    assert isinstance(result.error, OSError)

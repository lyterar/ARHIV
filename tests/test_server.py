"""Tests for the JSON API and ticket text generation."""

import json
import threading
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer

import pytest
from openpyxl import load_workbook

from server import make_handler
from tickets import ticket_text


@pytest.fixture
def api(tmp_path):
    path = tmp_path / "j.xlsx"
    server = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(path))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_address[1]}", path
    server.shutdown()


def post(url, body):
    req = urllib.request.Request(url, json.dumps(body).encode(), {"Content-Type": "application/json"})
    return urllib.request.urlopen(req)


def test_ticket_text_has_number_student_and_questions() -> None:
    text = ticket_text(7, "Иванов", "Иван")
    assert "№ 7" in text and "Иванов Иван" in text
    assert "1. " in text and "2. " in text


def test_every_ticket_has_text_and_out_of_range_fails() -> None:
    assert all(ticket_text(n) for n in range(1, 21))
    with pytest.raises(ValueError):
        ticket_text(21)


def test_post_registers_student_and_returns_text(api) -> None:
    url, path = api
    with post(url + "/api/tickets", {"last_name": "Иванов", "first_name": "Иван"}) as resp:
        data = json.load(resp)
        assert resp.status == 201
    assert 1 <= data["ticket"] <= 20 and "Иванов Иван" in data["text"]
    assert load_workbook(path).active.max_row == 2


def test_post_rejects_blank_names(api) -> None:
    with pytest.raises(urllib.error.HTTPError) as err:
        post(api[0] + "/api/tickets", {"last_name": " ", "first_name": "x"})
    assert err.value.code == 400


def test_get_ticket_text(api) -> None:
    with urllib.request.urlopen(api[0] + "/api/tickets/5") as resp:
        assert json.load(resp)["ticket"] == 5
    with pytest.raises(urllib.error.HTTPError) as err:
        urllib.request.urlopen(api[0] + "/api/tickets/99")
    assert err.value.code == 404

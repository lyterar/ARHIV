"""Tests for the console loop, using a fake console instead of the keyboard."""

from __future__ import annotations

import random
from pathlib import Path
from typing import Any, Iterable

import pytest
import readchar
from openpyxl import load_workbook

import journal
import main
from journal import HEADER
from main import (
    EMPTY_INPUT_MESSAGE,
    EXIT_HINT,
    FIRST_NAME_PROMPT,
    GOODBYE_MESSAGE,
    LAST_NAME_PROMPT,
    ExitRequested,
    KeyboardConsole,
    ask_name,
    run,
)
from tickets import generate_ticket


class FakeConsole:
    """Replays prepared lines, then behaves as if ESC was pressed."""

    def __init__(self, lines: Iterable[str], esc_on_retry: bool = False) -> None:
        self.lines = list(lines)
        self.prompts: list[str] = []
        self.retry_waits = 0
        self.esc_on_retry = esc_on_retry

    def read_line(self, prompt: str) -> str:
        self.prompts.append(prompt)
        if not self.lines:
            raise ExitRequested
        return self.lines.pop(0)

    def wait_for_retry(self) -> None:
        self.retry_waits += 1
        if self.esc_on_retry:
            raise ExitRequested


def fake_keys(monkeypatch: pytest.MonkeyPatch, *keys: str) -> None:
    """Make readchar.readkey() return *keys* one by one."""
    pending = iter(keys)
    monkeypatch.setattr(readchar, "readkey", lambda: next(pending))


def journal_rows(path: Path) -> list[tuple[Any, ...]]:
    return [tuple(r) for r in load_workbook(path).worksheets[0].iter_rows(values_only=True)]


# --- prompt loop -----------------------------------------------------------------


def test_blank_input_shows_message_and_repeats_same_prompt(
    capsys: pytest.CaptureFixture[str],
) -> None:
    console = FakeConsole(["", "   ", "  Иванов  "])

    assert ask_name(console, LAST_NAME_PROMPT) == "Иванов"
    assert console.prompts == [LAST_NAME_PROMPT] * 3
    assert capsys.readouterr().out.count(EMPTY_INPUT_MESSAGE) == 2


def test_run_registers_students_until_esc(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "journal.xlsx"
    console = FakeConsole([" Иванов ", "Иван", "Петрова", "", " Мария"])
    expected_rng = random.Random(5)
    tickets = [generate_ticket(expected_rng), generate_ticket(expected_rng)]

    run(console, path, random.Random(5))

    out = capsys.readouterr().out.splitlines()
    assert out[0] == EXIT_HINT
    assert f"Билет № {tickets[0]}" in out and f"Билет № {tickets[1]}" in out
    assert out[-1] == GOODBYE_MESSAGE
    rows = journal_rows(path)
    assert rows[0] == HEADER
    assert [row[:3] for row in rows[1:]] == [
        ("Иванов", "Иван", tickets[0]),
        ("Петрова", "Мария", tickets[1]),
    ]
    assert console.prompts[:5] == [
        LAST_NAME_PROMPT,
        FIRST_NAME_PROMPT,
        LAST_NAME_PROMPT,
        FIRST_NAME_PROMPT,
        FIRST_NAME_PROMPT,
    ]


def test_locked_journal_is_retried_with_the_same_record(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "journal.xlsx"
    real_replace = journal.os.replace
    calls = {"count": 0}

    def replace_locked_once(src: Any, dst: Any) -> None:
        calls["count"] += 1
        if calls["count"] == 1:
            raise PermissionError(13, "locked by Excel")
        real_replace(src, dst)

    monkeypatch.setattr(journal.os, "replace", replace_locked_once)
    console = FakeConsole(["Иванов", "Иван"])

    run(console, path, random.Random(1))

    assert console.retry_waits == 1
    assert (
        "Файл journal.xlsx открыт в другой программе. "
        "Закройте его и нажмите Enter для повторной попытки."
    ) in capsys.readouterr().out
    rows = journal_rows(path)
    assert len(rows) == 2 and rows[1][:2] == ("Иванов", "Иван")


def test_esc_while_waiting_for_retry_exits_and_reports_unsaved_record(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "journal.xlsx"
    path.write_bytes(b"broken")
    console = FakeConsole(["Иванов", "Иван"], esc_on_retry=True)

    run(console, path, random.Random(1))

    out = capsys.readouterr().out
    assert "повреждён" in out
    assert "Запись не сохранена: Иванов Иван" in out
    assert out.rstrip().endswith(GOODBYE_MESSAGE)
    assert path.read_bytes() == b"broken"


# --- keyboard handling -----------------------------------------------------------


def test_esc_as_first_key_requests_exit(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_keys(monkeypatch, "\x1b")
    with pytest.raises(ExitRequested):
        KeyboardConsole(prefill=lambda _p, _c: "unused").read_line(LAST_NAME_PROMPT)


def test_posix_esc_followed_by_another_key_requests_exit(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_keys(monkeypatch, "\x1bq")
    with pytest.raises(ExitRequested):
        KeyboardConsole(prefill=lambda _p, _c: "unused").read_line(LAST_NAME_PROMPT)


@pytest.mark.parametrize("enter", ["\r", "\n"])
def test_enter_as_first_key_gives_empty_line(monkeypatch: pytest.MonkeyPatch, enter: str) -> None:
    fake_keys(monkeypatch, enter)
    assert KeyboardConsole(prefill=lambda _p, _c: "unused").read_line(LAST_NAME_PROMPT) == ""


def test_first_character_is_handed_to_line_input(monkeypatch: pytest.MonkeyPatch) -> None:
    # Backspace with nothing typed, an arrow key and Tab are ignored.
    fake_keys(monkeypatch, "\x08", "\x7f", "\x00H", "\t", "И")
    seen: list[tuple[str, str]] = []

    def prefill(prompt: str, first: str) -> str:
        seen.append((prompt, first))
        return first + "ванов"

    assert KeyboardConsole(prefill).read_line(LAST_NAME_PROMPT) == "Иванов"
    assert seen == [(LAST_NAME_PROMPT, "И")]


def test_queued_first_character_is_left_for_line_input(monkeypatch: pytest.MonkeyPatch) -> None:
    # Windows path: the peek reports a typed character, so readchar is not used
    # and input() reads the whole line, first character included.
    fake_keys(monkeypatch)  # any readkey() call would raise StopIteration
    monkeypatch.setattr("builtins.input", lambda _prompt="": "Иванов")
    console = KeyboardConsole(prefill=lambda _p, _c: "unused", character_pending=lambda: True)

    assert console.read_line(LAST_NAME_PROMPT) == "Иванов"


def test_echo_prefill_joins_first_character_and_rest(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("builtins.input", lambda _prompt="": "ванов")
    assert main.echo_prefill(LAST_NAME_PROMPT, "И") == "Иванов"


def test_wait_for_retry_ignores_keys_until_enter(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_keys(monkeypatch, "a", " ", "\r")
    KeyboardConsole(main.echo_prefill).wait_for_retry()


def test_esc_while_waiting_for_retry_requests_exit(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_keys(monkeypatch, "a", "\x1b")
    with pytest.raises(ExitRequested):
        KeyboardConsole(main.echo_prefill).wait_for_retry()

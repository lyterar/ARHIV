"""Excel journal storage: create the workbook, append records, report file errors.

This module does no console I/O, so it can be tested on its own.
"""

from __future__ import annotations

import contextlib
import os
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

JOURNAL_FILENAME = "journal.xlsx"
HEADER = ("Last name", "First name", "Номер билета", "Дата и время")
COLUMN_WIDTHS = (20, 20, 15, 21)
EXCEL_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
DATETIME_COLUMN = 4


def clean_name(value: str) -> str | None:
    """Return *value* without control characters and surrounding whitespace.

    Returns None when nothing is left, i.e. the input counts as empty.
    """
    printable = "".join(ch for ch in value if ch.isprintable())
    return printable.strip() or None


@dataclass(frozen=True)
class Record:
    """One journal row. Names are cleaned on creation and must not be empty."""

    last_name: str
    first_name: str
    ticket: int
    timestamp: datetime = field(default_factory=datetime.now)

    def __post_init__(self) -> None:
        for attr in ("last_name", "first_name"):
            cleaned = clean_name(getattr(self, attr))
            if cleaned is None:
                raise ValueError(f"{attr} must not be empty")
            object.__setattr__(self, attr, cleaned)
        # The journal has one-second resolution; keep the stored value exact.
        object.__setattr__(self, "timestamp", self.timestamp.replace(microsecond=0))

    def as_row(self) -> tuple[str, str, int, datetime]:
        """Return the values in column order A..D."""
        return (self.last_name, self.first_name, self.ticket, self.timestamp)


class SaveStatus(Enum):
    """Outcome of an attempt to append a record."""

    OK = "ok"
    LOCKED = "locked"  # PermissionError: the file is open in another program
    CORRUPTED = "corrupted"  # the existing file is not a readable .xlsx workbook
    IO_ERROR = "io_error"  # any other OS-level failure (disk full, bad path, ...)


@dataclass(frozen=True)
class SaveResult:
    """Status of a save attempt plus the underlying exception, if any."""

    status: SaveStatus
    error: Exception | None = None

    @property
    def ok(self) -> bool:
        """True if the record was written to disk."""
        return self.status is SaveStatus.OK


def append_record(path: str | Path, record: Record) -> SaveResult:
    """Open (or create) the journal, append *record* after the last row, save, close.

    File problems never raise: they are returned as a SaveResult so the caller
    can keep the record and try again. Existing rows are never modified.
    """
    path = Path(path)
    try:
        workbook, sheet = _open_journal(path)
    except PermissionError as exc:
        return SaveResult(SaveStatus.LOCKED, exc)
    except OSError as exc:
        return SaveResult(SaveStatus.IO_ERROR, exc)
    except Exception as exc:  # zipfile, XML and openpyxl errors all mean "unreadable"
        return SaveResult(SaveStatus.CORRUPTED, exc)

    try:
        last_row = _last_data_row(sheet)
        if last_row == 0:
            _write_header(sheet)
            last_row = 1
        _write_record(sheet, last_row + 1, record)
        _save_atomically(workbook, path)
    except PermissionError as exc:
        return SaveResult(SaveStatus.LOCKED, exc)
    except OSError as exc:
        return SaveResult(SaveStatus.IO_ERROR, exc)
    finally:
        workbook.close()
    return SaveResult(SaveStatus.OK)


def _open_journal(path: Path) -> tuple[Workbook, Worksheet]:
    """Load the existing workbook or start a new one; return it with its first sheet."""
    workbook = load_workbook(path) if path.exists() else Workbook()
    return workbook, workbook.worksheets[0]


def _last_data_row(sheet: Worksheet) -> int:
    """Return the last row that holds any value, or 0 for an empty sheet.

    Rows that only carry formatting are ignored, so records stay contiguous.
    """
    for row_idx in range(sheet.max_row, 0, -1):
        if any(cell.value is not None for cell in sheet[row_idx]):
            return row_idx
    return 0


def _write_header(sheet: Worksheet) -> None:
    """Fill row 1 with bold column titles and set readable column widths."""
    bold = Font(bold=True)
    for col_idx, (title, width) in enumerate(zip(HEADER, COLUMN_WIDTHS), start=1):
        sheet.cell(row=1, column=col_idx, value=title).font = bold
        sheet.column_dimensions[get_column_letter(col_idx)].width = width


def _write_record(sheet: Worksheet, row_idx: int, record: Record) -> None:
    """Write *record* into row *row_idx*."""
    for col_idx, value in enumerate(record.as_row(), start=1):
        cell = sheet.cell(row=row_idx, column=col_idx, value=value)
        if isinstance(value, str):
            cell.data_type = "s"  # a name like "=1+1" must stay text, not a formula
    sheet.cell(row=row_idx, column=DATETIME_COLUMN).number_format = EXCEL_DATETIME_FORMAT


def _save_atomically(workbook: Workbook, path: Path) -> None:
    """Write to a temp file next to *path*, flush it to disk, then swap it in.

    If the process dies mid-save, the previous journal stays intact instead of
    being left truncated.
    """
    tmp_path = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with open(tmp_path, "wb") as tmp_file:
            workbook.save(tmp_file)
            tmp_file.flush()
            os.fsync(tmp_file.fileno())
        os.replace(tmp_path, path)
    finally:
        with contextlib.suppress(OSError):
            tmp_path.unlink(missing_ok=True)  # already gone after a successful replace
